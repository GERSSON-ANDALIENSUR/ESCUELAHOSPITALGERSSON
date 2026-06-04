#!/usr/bin/env python3
"""Genera el Bingo de Pueblos Originarios de Chile - 24 cartones."""

import random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter

random.seed(42)

# ─────────────────────────────────────────────────────────────
# DATOS
# ─────────────────────────────────────────────────────────────
ELEMENTS = [
    "Mapuche", "Aymara", "Rapa Nui", "Quechua", "Atacameño", "Colla",
    "Diaguita", "Chango", "Kawésqar", "Yagán", "Selk'nam", "Moai",
    "Kultrún", "Trarilonco", "Vasija Diaguita", "Máscara Selk'nam",
    "Ruca", "Canoa", "Totora", "Alpaca", "Zampoña", "Cesta",
    "Quena", "Desierto",
]

CATEGORIES = [
    "Pueblo originario", "Pueblo originario", "Pueblo originario",
    "Pueblo originario", "Pueblo originario", "Pueblo originario",
    "Pueblo originario", "Pueblo originario", "Pueblo originario",
    "Pueblo originario", "Pueblo originario", "Objeto cultural",
    "Instrumento musical", "Objeto cultural", "Objeto cultural",
    "Objeto cultural", "Vivienda o construcción", "Objeto cultural",
    "Elemento natural o territorial", "Animal",
    "Instrumento musical", "Objeto cultural",
    "Instrumento musical", "Elemento natural o territorial",
]

PROMPTS = [
    # 1. Mapuche
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Mapuche, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 2. Aymara
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Aymara, vestimenta cultural sencilla con colores andinos. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 3. Rapa Nui
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Rapa Nui, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 4. Quechua
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Quechua, vestimenta cultural sencilla con colores andinos. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 5. Atacameño
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Atacameño, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 6. Colla
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Colla, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 7. Diaguita
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Diaguita, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 8. Chango
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Chango, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 9. Kawésqar
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Kawésqar, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 10. Yagán
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Yagán, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 11. Selk'nam
    "Imagen cuadrada, formato LEGO animado 3D. Una figura LEGO que representa respetuosamente a una persona del pueblo Selk'nam, vestimenta cultural sencilla. Fondo blanco liso, sin decoraciones, sin paisaje, sin objetos adicionales. Solo la figura, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 12. Moai
    "Imagen cuadrada, formato LEGO animado 3D. Un solo moai de Rapa Nui construido con bloques LEGO, visto de frente. Fondo blanco liso, sin hierba, sin cielo, sin decoraciones adicionales. Solo el moai, centrado. Sin texto, sin logos, sin marcas de agua.",
    # 13. Kultrún
    "Imagen cuadrada, formato LEGO animado 3D. Un solo kultrún mapuche, tambor circular con sus diseños geométricos visibles, visto desde arriba o de frente. Fondo blanco liso, sin mesa, sin decoraciones adicionales. Solo el kultrún, centrado. Sin texto, sin logos, sin marcas de agua.",
    # 14. Trarilonco
    "Imagen cuadrada, formato LEGO animado 3D. Un solo trarilonco mapuche, vincha ornamental de plata, flotando sobre fondo blanco liso. Sin vitrina, sin mesa, sin decoraciones adicionales. Solo el trarilonco, centrado. Sin texto, sin logos, sin marcas de agua.",
    # 15. Vasija Diaguita
    "Imagen cuadrada, formato LEGO animado 3D. Una sola vasija diaguita con diseños geométricos en rojo y negro, vista de frente. Fondo blanco liso, sin mesa, sin decoraciones adicionales. Solo la vasija, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 16. Máscara Selk'nam
    "Imagen cuadrada, formato LEGO animado 3D. Una sola máscara Selk'nam con colores naturales, vista de frente, representada con respeto como objeto cultural. Fondo blanco liso, sin decoraciones adicionales. Solo la máscara, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 17. Ruca
    "Imagen cuadrada, formato LEGO animado 3D. Una sola ruca mapuche construida con bloques LEGO, vista de frente o en perspectiva ligera. Fondo blanco liso, sin árboles, sin pasto, sin decoraciones adicionales. Solo la ruca, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 18. Canoa
    "Imagen cuadrada, formato LEGO animado 3D. Una sola canoa tradicional de pueblos canoeros del extremo sur de Chile, construida con bloques LEGO, vista de costado. Fondo blanco liso, sin agua, sin montañas, sin decoraciones adicionales. Solo la canoa, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 19. Totora
    "Imagen cuadrada, formato LEGO animado 3D. Una sola mata de totora, planta alta con tallos verdes y espiga marrón, construida con bloques LEGO. Fondo blanco liso, sin lago, sin montañas, sin decoraciones adicionales. Solo la planta de totora, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 20. Alpaca
    "Imagen cuadrada, formato LEGO animado 3D. Una sola alpaca andina construida con bloques LEGO, vista de frente o de costado. Fondo blanco liso, sin montañas, sin pasto, sin decoraciones adicionales. Solo la alpaca, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 21. Zampoña
    "Imagen cuadrada, formato LEGO animado 3D. Una sola zampoña andina, instrumento de tubos de caña, construida con bloques LEGO, vista de frente. Fondo blanco liso, sin manta, sin montañas, sin decoraciones adicionales. Solo la zampoña, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 22. Cesta
    "Imagen cuadrada, formato LEGO animado 3D. Una sola cesta artesanal tejida con fibras naturales, construida con bloques LEGO, vista ligeramente desde arriba para mostrar el tejido. Fondo blanco liso, sin decoraciones adicionales. Solo la cesta, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 23. Quena
    "Imagen cuadrada, formato LEGO animado 3D. Una sola quena, flauta andina de caña, construida con bloques LEGO, vista de frente. Fondo blanco liso, sin manta, sin montañas, sin decoraciones adicionales. Solo la quena, centrada. Sin texto, sin logos, sin marcas de agua.",
    # 24. Desierto
    "Imagen cuadrada, formato LEGO animado 3D. Una escena mínima del desierto de Atacama: solo arena y una pequeña duna, construida con bloques LEGO, sin personajes ni objetos adicionales. Fondo celeste muy liso, sin montañas, sin vegetación, sin decoraciones. Solo el desierto, centrado. Sin texto, sin logos, sin marcas de agua.",
]

# ─────────────────────────────────────────────────────────────
# GENERACIÓN DE CARTONES (diseño circulante equilibrado)
# Cada elemento aparece exactamente 12 veces en los 24 cartones.
# ─────────────────────────────────────────────────────────────
def generate_cards():
    cards = []
    for i in range(24):
        card = [ELEMENTS[(i + j) % 24] for j in range(12)]
        random.shuffle(card)
        cards.append(card)
    return cards


# ─────────────────────────────────────────────────────────────
# UTILIDADES DE ESTILO
# ─────────────────────────────────────────────────────────────
def border(style="medium"):
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)


def cell_style(ws, row, col, value, *, font_size=12, bold=False,
               bg_color=None, font_color="2C1810", h_align="center",
               wrap=True, border_style="medium"):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    cell.font = Font(name="Calibri", size=font_size, bold=bold, color=font_color)
    cell.alignment = Alignment(horizontal=h_align, vertical="center", wrap_text=wrap)
    if bg_color:
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color,
                                fill_type="solid")
    cell.border = border(border_style)
    return cell


# ─────────────────────────────────────────────────────────────
# HOJA 1 – CARTONES BINGO
# ─────────────────────────────────────────────────────────────
# Paleta: (fondo casilla, fondo título, color texto título)
SCHEMES = [
    ("FFF3E0", "E65100", "FFFFFF"),   # Naranja tierra
    ("E8F5E9", "2E7D32", "FFFFFF"),   # Verde naturaleza
    ("E3F2FD", "1565C0", "FFFFFF"),   # Azul cielo
    ("FFFDE7", "F57F17", "FFFFFF"),   # Amarillo sol
    ("F3E5F5", "6A1B9A", "FFFFFF"),   # Violeta suave
    ("FBE9E7", "BF360C", "FFFFFF"),   # Terracota
]


def build_sheet1(wb, cards):
    ws = wb.active
    ws.title = "Cartones Bingo"

    COLS_PER_CARD = 4
    DATA_ROWS = 3
    CARDS_PER_ROW = 3
    COL_GAP = 2      # columnas vacías entre cartones
    ROW_GAP = 1      # fila vacía entre filas de cartones
    BLOCK_ROWS = 1 + DATA_ROWS + ROW_GAP  # 5 filas por bloque

    for idx, card in enumerate(cards):
        c_row = idx // CARDS_PER_ROW
        c_col = idx % CARDS_PER_ROW

        sr = c_row * BLOCK_ROWS + 1                     # fila inicio
        sc = c_col * (COLS_PER_CARD + COL_GAP) + 1     # col inicio

        bg, title_bg, title_fg = SCHEMES[idx % len(SCHEMES)]

        # ── Título (celda combinada) ──────────────────────────
        title_cell = ws.cell(row=sr, column=sc)
        title_cell.value = f"Cartón {idx + 1}"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color=title_fg)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.fill = PatternFill(start_color=title_bg, end_color=title_bg,
                                      fill_type="solid")

        # Borde exterior de la fila título (celdas individuales antes de combinar)
        med = Side(style="medium")
        thin = Side(style="thin")
        for c_off in range(COLS_PER_CARD):
            tc = ws.cell(row=sr, column=sc + c_off)
            is_first = c_off == 0
            is_last = c_off == COLS_PER_CARD - 1
            tc.border = Border(
                left=med if is_first else Side(style=None),
                right=med if is_last else Side(style=None),
                top=med,
                bottom=thin,
            )
            tc.fill = PatternFill(start_color=title_bg, end_color=title_bg,
                                  fill_type="solid")

        ws.merge_cells(start_row=sr, start_column=sc,
                       end_row=sr, end_column=sc + COLS_PER_CARD - 1)

        # ── Casillas de datos ─────────────────────────────────
        for r in range(DATA_ROWS):
            for c in range(COLS_PER_CARD):
                elem = card[r * COLS_PER_CARD + c]
                dcell = ws.cell(row=sr + 1 + r, column=sc + c)
                dcell.value = elem
                dcell.font = Font(name="Calibri", size=12, bold=False, color="2C1810")
                dcell.alignment = Alignment(horizontal="center", vertical="center",
                                            wrap_text=True)
                dcell.fill = PatternFill(start_color=bg, end_color=bg,
                                         fill_type="solid")

                is_top = r == 0
                is_bottom = r == DATA_ROWS - 1
                is_left = c == 0
                is_right = c == COLS_PER_CARD - 1
                dcell.border = Border(
                    left=med if is_left else thin,
                    right=med if is_right else thin,
                    top=thin if is_top else thin,
                    bottom=med if is_bottom else thin,
                )

    # ── Anchos de columna ──────────────────────────────────────
    for c_col in range(CARDS_PER_ROW):
        sc = c_col * (COLS_PER_CARD + COL_GAP) + 1
        for off in range(COLS_PER_CARD):
            ws.column_dimensions[get_column_letter(sc + off)].width = 18
        for g in range(COL_GAP):
            ws.column_dimensions[get_column_letter(sc + COLS_PER_CARD + g)].width = 3

    # ── Altos de fila ─────────────────────────────────────────
    total_card_rows = (24 + CARDS_PER_ROW - 1) // CARDS_PER_ROW
    for c_row in range(total_card_rows):
        sr = c_row * BLOCK_ROWS + 1
        ws.row_dimensions[sr].height = 26        # título
        for r in range(1, DATA_ROWS + 1):
            ws.row_dimensions[sr + r].height = 42  # casillas
        ws.row_dimensions[sr + DATA_ROWS + 1].height = 10  # separador

    # ── Configuración de impresión ───────────────────────────
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True


# ─────────────────────────────────────────────────────────────
# HOJA 2 – PROMPTS GEMINI
# ─────────────────────────────────────────────────────────────
def build_sheet2(wb):
    ws = wb.create_sheet("Prompts Gemini")

    headers = ["N°", "Elemento", "Categoría", "Prompt para Gemini"]
    header_bg = "1B5E20"
    alt_colors = ["FFFFFF", "F1F8E9"]

    # ── Encabezado ────────────────────────────────────────────
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.value = h
        c.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = PatternFill(start_color=header_bg, end_color=header_bg,
                             fill_type="solid")
        c.border = border("medium")
    ws.row_dimensions[1].height = 30

    # ── Filas de datos ────────────────────────────────────────
    for i in range(24):
        row = i + 2
        bg = alt_colors[i % 2]
        row_data = [i + 1, ELEMENTS[i], CATEGORIES[i], PROMPTS[i]]
        aligns = ["center", "center", "center", "left"]

        for col, (val, align) in enumerate(zip(row_data, aligns), 1):
            c = ws.cell(row=row, column=col)
            c.value = val
            c.font = Font(name="Calibri", size=11)
            c.alignment = Alignment(horizontal=align, vertical="center",
                                    wrap_text=True)
            c.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            c.border = border("thin")

        ws.row_dimensions[row].height = 72

    # ── Anchos de columna ──────────────────────────────────────
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 88

    # ── Impresión ─────────────────────────────────────────────
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1


# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────
def main():
    cards = generate_cards()

    # Verificaciones
    assert len(cards) == 24, "Número de cartones incorrecto"
    assert all(len(c) == 12 for c in cards), "Algún cartón no tiene 12 elementos"
    counts = Counter(e for c in cards for e in c)
    assert all(v == 12 for v in counts.values()), f"Distribución no equilibrada: {counts}"
    # Verificar que no haya dos cartones iguales (como conjuntos)
    sets = [frozenset(c) for c in cards]
    assert len(set(sets)) == 24, "Hay cartones duplicados"

    wb = Workbook()
    build_sheet1(wb, cards)
    build_sheet2(wb)

    output = "/home/user/ESCUELAHOSPITALGERSSON/Bingo_Pueblos_Originarios_24_Cartones.xlsx"
    wb.save(output)
    print(f"✓ Archivo guardado: {output}")

    print("\nDistribución de elementos (cada uno debe aparecer 12 veces):")
    for elem in ELEMENTS:
        print(f"  {counts[elem]:2d}×  {elem}")

    print(f"\n✓ 24 cartones generados")
    print(f"✓ 12 casillas por cartón")
    print(f"✓ 24 prompts en hoja 2")
    print(f"✓ Sin cartones duplicados")


if __name__ == "__main__":
    main()
